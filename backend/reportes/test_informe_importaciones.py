from django.test import TestCase, override_settings

from comercio.busqueda import texto_producto_importacion
from comercio.models import ArchivoCarga, Importacion
from reportes.models import ImportadorProbable
from reportes.views import DIN_LABELS, _filtered_importaciones


def din_row(**positions):
    raw = [""] * len(DIN_LABELS)
    for index, value in positions.items():
        raw[int(index.lstrip("p"))] = value
    return {"raw_columns": raw}


@override_settings(ALLOWED_HOSTS=["testserver"])
class InformeImportacionesUniversoTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        source = ArchivoCarga.objects.create(nombre_archivo="x", archivo="cargas/x.zip", tipo_archivo="IMP")
        cls.sony = ImportadorProbable.objects.create(rut="76123456", dv="0", nombre="SONY CHILE LTDA")
        common = {"archivo_origen": source, "periodo_anio": 2026, "periodo_mes": 2}
        # DIN TXT: brand in DMARCA (134), description split mid-word across DOTRO1/DOTRO2 (136/137).
        cls.tv = Importacion.objects.create(
            **common, numero_ident="1", aduana_codigo="39", pais_origen_codigo="336", importador_probable_sugerido=cls.sony,
            payload_json=din_row(p133="TELEVISOR", p134="SONY-F", p136="PANTALLA LED CON T", p137="ODOS SUS ACCESORIOS", p22="225"),
        )
        # MDB 2015 rows keep named fields instead of DIN positions.
        cls.mango = Importacion.objects.create(
            **common, numero_ident="2", aduana_codigo="14", pais_origen_codigo="219",
            payload_json={"legacy_fields": {"MERCADERIA": "MANGO", "ATRI1": "MACHU PICCHU-F", "PAI_ADQ": "219"}},
        )

        for row in Importacion.objects.all():
            row.texto_producto = texto_producto_importacion(row.payload_json)
            row.save(update_fields=["texto_producto"])

    def ids(self, filters):
        return set(_filtered_importaciones(filters, 2026, 2).values_list("numero_ident", flat=True))

    def test_product_terms_search_all_description_fields_and_add_up(self):
        self.assertEqual(self.ids({"productos": ["sony"]}), {"1"})
        self.assertEqual(self.ids({"productos": ["machu picchu"]}), {"2"})
        self.assertEqual(self.ids({"productos": ["sony", "mango"]}), {"1", "2"})
        self.assertEqual(self.ids({"productos": ["todos sus accesorios"]}), {"1"})

    def test_country_of_acquisition_reads_din_position_and_mdb_field(self):
        self.assertEqual(self.ids({"pais_adquisicion_codigo": ["225"]}), {"1"})
        self.assertEqual(self.ids({"pais_adquisicion_codigo": ["219"]}), {"2"})

    def test_importer_origin_country_and_customs_filters(self):
        self.assertEqual(self.ids({"importadores": [str(self.sony.id)]}), {"1"})
        self.assertEqual(self.ids({"pais_origen_codigo": ["219"], "aduana_codigo": ["14"]}), {"2"})
        self.assertEqual(self.ids({"pais_origen_codigo": ["336"], "aduana_codigo": ["14"]}), set())

    def test_importer_search_by_name_rut_and_ids(self):
        by_name = self.client.get("/api/reportes/importadores/", {"q": "sony"}).json()
        by_rut = self.client.get("/api/reportes/importadores/", {"q": "76.123"}).json()
        by_ids = self.client.get("/api/reportes/importadores/", {"ids": str(self.sony.id)}).json()
        self.assertEqual([row["nombre"] for row in by_name + by_rut + by_ids], ["SONY CHILE LTDA"] * 3)


@override_settings(ALLOWED_HOSTS=["testserver"], CELERY_TASK_ALWAYS_EAGER=True)
class InformeImportacionesRangoTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        source = ArchivoCarga.objects.create(nombre_archivo="x", archivo="cargas/x.zip", tipo_archivo="IMP")
        for anio, mes in ((2015, 6), (2025, 12), (2026, 1), (2026, 2), (2026, 3)):
            Importacion.objects.create(archivo_origen=source, periodo_anio=anio, periodo_mes=mes, numero_ident=f"{anio}-{mes}", payload_json=din_row())
        from django.contrib.auth import get_user_model

        cls.user = get_user_model().objects.create_user("analista", password="x")
        cls.other = get_user_model().objects.create_user("otro", password="x")

    def setUp(self):
        import tempfile
        from unittest.mock import patch

        self.client.force_login(self.user)
        self.media = tempfile.TemporaryDirectory()
        patcher = patch("reportes.tasks.INFORMES_DIR", __import__("pathlib").Path(self.media.name))
        patcher.start()
        self.addCleanup(patcher.stop)
        self.addCleanup(self.media.cleanup)

    def post(self, desde, hasta):
        import json

        body = {"columnas": ["numero_ident"], "filtros": {}, "periodo_desde": {"anio": desde[0], "mes": desde[1]}, "periodo_hasta": {"anio": hasta[0], "mes": hasta[1]}}
        return self.client.post("/api/reportes/importaciones/exportar/", data=json.dumps(body), content_type="application/json")

    def test_range_crosses_years_and_is_inclusive(self):
        ids = set(_filtered_importaciones({}, desde=(2025, 12), hasta=(2026, 2)).values_list("numero_ident", flat=True))
        self.assertEqual(ids, {"2025-12", "2026-1", "2026-2"})

    def test_generates_in_background_and_downloads_only_for_owner(self):
        import io

        from openpyxl import load_workbook

        response = self.post((2026, 1), (2026, 3))
        self.assertEqual(response.status_code, 202)
        informe = response.json()
        self.assertEqual((informe["estado"], informe["filas_total"], informe["nombre_descarga"]), ("LISTO", 3, "informe_importaciones_2026-01_a_2026-03.xlsx"))
        download = self.client.get(f"/api/reportes/informes/{informe['id']}/descargar/")
        rows = list(load_workbook(io.BytesIO(b"".join(download.streaming_content)), read_only=True).active.iter_rows(values_only=True))
        self.assertEqual(sorted(row[0] for row in rows[1:]), ["2026-1", "2026-2", "2026-3"])
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(f"/api/reportes/informes/{informe['id']}/").status_code, 404)

    def test_rejects_inverted_range_and_reports_excel_row_limit(self):
        from unittest.mock import patch

        self.assertEqual(self.post((2026, 3), (2026, 1)).status_code, 400)
        with patch("reportes.tasks.EXCEL_MAX_ROWS", 2):
            informe = self.post((2015, 1), (2026, 12)).json()
        self.assertEqual(informe["estado"], "ERROR")
        self.assertIn("Excel admite", informe["error"])


class ProductoSinTildesTests(TestCase):
    def test_accented_term_matches_unaccented_customs_text(self):
        source = ArchivoCarga.objects.create(nombre_archivo="x", archivo="cargas/x.zip", tipo_archivo="IMP")
        payload = {"legacy_fields": {"MERCADERIA": "NEUMATICOS RADIALES"}}
        Importacion.objects.create(archivo_origen=source, periodo_anio=2015, periodo_mes=1, numero_ident="1", payload_json=payload, texto_producto=texto_producto_importacion(payload))
        self.assertEqual(_filtered_importaciones({"productos": ["neumáticos"]}, 2015, 1).count(), 1)
