import io
import json

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from openpyxl import load_workbook

from catalogos.models import CatalogoCodigo, PartidaArancelaria
from comercio.export_loader import export_members, load_member
from comercio.test_exportaciones import BASE, BULTOS, DOCS, dus_line, make_zip


@override_settings(ALLOWED_HOSTS=["testserver"])
class ExportacionesConsultaInformeTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        archive = make_zip({
            BASE: [dus_line(), dus_line(item="2", FOBUS="100", CODIGOARANCEL="08061099", NOMBRE="UVAS FRESCAS", ATRIBUTO1="RED GLOBE")],
            BULTOS: ["13213653;24022026;1;22;5;", "13213653;24022026;2;22;7;"],
            DOCS: ["13213653;24022026;1;MAEU1;10012026;POLAR ARGENTINA;266103"],
        })
        for member, tipo in export_members(archive):
            load_member(archive, member, tipo, archivo_path="cargas/x.zip")
        CatalogoCodigo.objects.create(grupo="paises", codigo="218", glosa="ECUADOR")
        CatalogoCodigo.objects.create(grupo="tipos_bulto", codigo="22", glosa="CAJA DE CARTON")
        PartidaArancelaria.objects.create(codigo="08061099", glosa="Las demás uvas frescas")
        cls.user = get_user_model().objects.create_user("analista", password="x")

    def setUp(self):
        self.client.force_login(self.user)

    def test_requires_session(self):
        self.client.logout()
        self.assertEqual(self.client.get("/api/consultas/exportaciones/").status_code, 403)
        self.assertEqual(self.client.get("/api/reportes/exportaciones/configuracion/").status_code, 403)

    def test_search_by_tariff_description_with_catalog_glosas(self):
        data = self.client.get("/api/consultas/exportaciones/", {"partida_busqueda": "uvas", "periodo_anio": "2026"}).json()
        self.assertEqual(data["count"], 1)
        row = data["results"][0]
        self.assertEqual((row["item"], row["pais_destino_glosa"], row["partida_glosa"]), ("2", "ECUADOR", "Las demás uvas frescas"))

    def test_dus_detail_includes_bultos_and_documents(self):
        data = self.client.get("/api/consultas/exportaciones/dus/13213653/").json()
        self.assertEqual(len(data["items"]), 2)
        self.assertEqual([bulto["tipo_bulto_glosa"] for bulto in data["bultos"]], ["CAJA DE CARTON", "CAJA DE CARTON"])
        self.assertEqual(data["documentos"][0]["nave"], "POLAR ARGENTINA")

    def report(self, filtros, columnas, desde=(2026, 2), hasta=(2026, 2)):
        """Request a report (generated eagerly in tests) and return its spreadsheet rows."""
        import tempfile
        from pathlib import Path
        from unittest.mock import patch

        body = {"columnas": columnas, "filtros": filtros, "periodo_desde": {"anio": desde[0], "mes": desde[1]}, "periodo_hasta": {"anio": hasta[0], "mes": hasta[1]}}
        with tempfile.TemporaryDirectory() as folder, patch("reportes.tasks.INFORMES_DIR", Path(folder)), override_settings(CELERY_TASK_ALWAYS_EAGER=True):
            informe = self.client.post("/api/reportes/exportaciones/exportar/", data=json.dumps(body), content_type="application/json").json()
            self.assertEqual(informe["estado"], "LISTO", informe)
            download = self.client.get(f"/api/reportes/informes/{informe['id']}/descargar/")
            return list(load_workbook(io.BytesIO(b"".join(download.streaming_content)), read_only=True).active.iter_rows(values_only=True))

    def test_report_writes_numbers_dus_fields_and_bulto_totals(self):
        config = self.client.get("/api/reportes/exportaciones/configuracion/").json()
        self.assertEqual(config["periodos"], [{"anio": 2026, "mes": 2}])
        rows = self.report(
            {"partidas": ["08061099"]},
            ["numero_ident", "pais_destino_glosa", "valor_fob", "dus:ATRIBUTO1", "total_bultos", "tipos_bulto", "naves"],
        )
        self.assertEqual(rows[1], ("13213653", "ECUADOR", 100, "RED GLOBE", 12, "CAJA DE CARTON", "POLAR ARGENTINA"))
        self.assertEqual(len(rows), 2)

    def test_report_filters_products_without_accents_exporters_and_range(self):
        self.assertEqual([row[0] for row in self.report({"productos": ["úvas", "red globe"]}, ["item"])[1:]], ["2"])
        self.assertEqual(len(self.report({"exportadores": ["7435"]}, ["item"])), 3)
        self.assertEqual(len(self.report({"exportadores": ["1"]}, ["item"])), 1)
        self.assertEqual(len(self.report({}, ["item"], desde=(2026, 3), hasta=(2026, 9))), 1)

    def test_export_rubros_are_separate_from_import_rubros(self):
        payload = json.dumps({"nombre": "Fruta", "configuracion_json": {"columnas": ["item"]}})
        self.assertEqual(self.client.post("/api/reportes/exportaciones/rubros/", data=payload, content_type="application/json").status_code, 201)
        self.assertEqual(self.client.post("/api/reportes/importaciones/rubros/", data=payload, content_type="application/json").status_code, 201)
        self.assertEqual(self.client.post("/api/reportes/exportaciones/rubros/", data=payload, content_type="application/json").status_code, 400)
        self.assertEqual([r["nombre"] for r in self.client.get("/api/reportes/exportaciones/rubros/").json()], ["Fruta"])

    def test_report_requires_period(self):
        response = self.client.post("/api/reportes/exportaciones/exportar/", data=json.dumps({"columnas": ["numero_ident"]}), content_type="application/json")
        self.assertEqual(response.status_code, 400)
